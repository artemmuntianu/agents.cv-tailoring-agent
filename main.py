import os
import sys
from agent.graph import create_graph
from agent.state import State

def run_cv_tailoring(cv_path: str, job_desc_path: str, output_path: str):
    print("🚀 Starting Automated CV Tailoring Agent...")
    print(f"📂 CV Path: {cv_path}")
    print(f"📋 Job Description Path: {job_desc_path}")
    print(f"🎯 Target Output Path: {output_path}")

    if not os.path.exists(cv_path):
        print(f"❌ Error: Resume file not found at {cv_path}")
        sys.exit(1)
        
    if not os.path.exists(job_desc_path):
        print(f"❌ Error: Job description file not found at {job_desc_path}")
        sys.exit(1)

    with open(job_desc_path, "r", encoding="utf-8") as f:
        job_description_text = f.read()

    initial_state: State = {
        "cv_path": cv_path,
        "job_description": job_description_text,
        "output_path": output_path,
        "current_cv_text": "",
        "modifications": [],
        "layout_feedback": "",
        "revision_count": 0,
        "image_paths": [],
        "is_approved": False
    }

    graph = create_graph()
    final_state = graph.invoke(initial_state)

    print("\n✨ Process finished!")
    print(f"📄 Tailored document saved to: {final_state['output_path']}")
    print(f"🏆 Final Layout Approved: {final_state['is_approved']}")
    print(f"🔄 Total Revisions: {final_state['revision_count']}")

if __name__ == "__main__":
    sample_cv = os.path.abspath("sample_cv.docx")
    sample_jd = os.path.abspath("job_description.txt")
    output_cv = os.path.abspath("output_tailored_cv.docx")

    if not os.path.exists(sample_cv) or not os.path.exists(sample_jd):
        print("💡 Creating sample input files for demonstration...")
        import docx
        doc = docx.Document()
        doc.add_heading("Jane Doe - Senior AI Engineer", level=0)
        p1 = doc.add_paragraph("Summary: Experienced Python Engineer building LLM applications and automated tools.")
        p2 = doc.add_paragraph("Experience: Built scalable microservices using Python and FastAPI. Implemented search pipelines with vector databases.")
        doc.save(sample_cv)

        with open(sample_jd, "w", encoding="utf-8") as f:
            f.write("We are looking for a Senior AI Engineer skilled in Python, LangGraph, Google GenAI SDK, and automated document processing workflows.")

    run_cv_tailoring(sample_cv, sample_jd, output_cv)
