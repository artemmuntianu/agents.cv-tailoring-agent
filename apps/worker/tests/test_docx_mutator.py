"""DOCX mutation behaviour (the part that must never regress).

The text rendering lives in `test_cv_text.py` and the replacement rules in
`test_cv_replacements.py`; this file covers the AST surgery and the sync validator.
"""

import os
import tempfile

import docx

from tests.helpers import SAMPLE_CV_DATA, docx_lines, write_docx
from utils import docx_mutator


def test_apply_text_replacements_rewrites_the_paragraph_and_saves():
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "master.docx")
        write_docx(docx_path, docx_lines(SAMPLE_CV_DATA))
        output_path = os.path.join(tmp, "out", "tailored.docx")

        applied = docx_mutator.apply_text_replacements(
            docx_path,
            [(SAMPLE_CV_DATA["summary"], "Platform Engineering Lead with Azure delivery record.")],
            output_path,
        )

        assert applied == 1
        assert os.path.exists(output_path)
        paragraphs = [p.text for p in docx.Document(output_path).paragraphs]
        assert "Platform Engineering Lead with Azure delivery record." in paragraphs
        # The master document itself must stay untouched.
        master_text = [p.text for p in docx.Document(docx_path).paragraphs]
        assert SAMPLE_CV_DATA["summary"] in master_text


def test_validate_cv_data_against_docx_detects_drift():
    with tempfile.TemporaryDirectory() as tmp:
        in_sync = os.path.join(tmp, "in_sync.docx")
        write_docx(in_sync, docx_lines(SAMPLE_CV_DATA))
        assert docx_mutator.validate_cv_data_against_docx(SAMPLE_CV_DATA, in_sync) == []

        drifted = os.path.join(tmp, "drifted.docx")
        write_docx(drifted, ["Completely different resume content."])
        missing = docx_mutator.validate_cv_data_against_docx(SAMPLE_CV_DATA, drifted)
        assert SAMPLE_CV_DATA["summary"] in missing
