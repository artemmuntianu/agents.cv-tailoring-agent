"""DOCX mutation behaviour (the part that must never regress)."""

import os
import tempfile

import docx

from tests.helpers import SAMPLE_CV_DATA, docx_lines, isolated_config, write_docx
from utils import docx_mutator


def test_cv_data_to_text_contains_every_section():
    text = docx_mutator.cv_data_to_text(SAMPLE_CV_DATA)
    assert "SUMMARY:" in text
    assert "RELEVANT SKILLS:" in text
    assert "PROFESSIONAL EXPERIENCE:" in text
    assert "Languages" in text
    assert "• Led the migration of 50 desktop screens to a web platform." in text


def test_extract_doc_text_accepts_cv_data_and_legacy_path():
    from_data = docx_mutator.extract_doc_text(SAMPLE_CV_DATA)
    assert from_data.startswith(SAMPLE_CV_DATA["header"]["title"])

    with tempfile.TemporaryDirectory() as tmp:
        with isolated_config(tmp):
            # Legacy call style: a bare path string used to be accepted.
            via_path = docx_mutator.extract_doc_text(os.path.join(tmp, "whatever.docx"))
    assert via_path == from_data


def test_normalize_replacements_splits_concatenated_label_and_value():
    items = [("Languages\nC#, SQL", "Languages (ATS)\nC#, SQL, Azure")]
    result = docx_mutator.normalize_replacements(items)
    assert result == [
        ("Languages", "Languages (ATS)", "N/A"),
        ("C#, SQL", "C#, SQL, Azure", "N/A"),
    ]


def test_normalize_replacements_drops_misaligned_lines_bullets_and_noops():
    items = [
        ("a\nb", "only-one-line"),  # cannot align -> dropped
        ("• Led the migration", "• Led the migration"),  # no-op -> dropped
        ("• Cut report generation time by 80%.", "Cut report time by 80%.", "ATS keyword"),
    ]
    result = docx_mutator.normalize_replacements(items)
    assert result == [("Cut report generation time by 80%.", "Cut report time by 80%.", "ATS keyword")]


def test_apply_text_replacements_rewrites_the_paragraph_and_saves():
    with tempfile.TemporaryDirectory() as tmp:
        docx_path = os.path.join(tmp, "master.docx")
        write_docx(docx_path, docx_lines(SAMPLE_CV_DATA))
        output_path = os.path.join(tmp, "out", "tailored.docx")

        applied = docx_mutator.apply_text_replacements(
            docx_path,
            [(SAMPLE_CV_DATA["summary"], "Platform Engineering Lead with Azure delivery record.")],
            output_path,
        )

        assert applied == 1
        assert os.path.exists(output_path)
        paragraphs = [p.text for p in docx.Document(output_path).paragraphs]
        assert "Platform Engineering Lead with Azure delivery record." in paragraphs
        # The master document itself must stay untouched.
        master_text = [p.text for p in docx.Document(docx_path).paragraphs]
        assert SAMPLE_CV_DATA["summary"] in master_text


def test_validate_cv_data_against_docx_detects_drift():
    with tempfile.TemporaryDirectory() as tmp:
        in_sync = os.path.join(tmp, "in_sync.docx")
        write_docx(in_sync, docx_lines(SAMPLE_CV_DATA))
        assert docx_mutator.validate_cv_data_against_docx(SAMPLE_CV_DATA, in_sync) == []

        drifted = os.path.join(tmp, "drifted.docx")
        write_docx(drifted, ["Completely different resume content."])
        missing = docx_mutator.validate_cv_data_against_docx(SAMPLE_CV_DATA, drifted)
        assert SAMPLE_CV_DATA["summary"] in missing
